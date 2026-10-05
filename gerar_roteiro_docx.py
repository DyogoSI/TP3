"""Gera roteiro_apresentacao.docx a partir das notas de slides_tp3.pptx."""
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pptx import Presentation

NAVY, GREEN, MUTED = RGBColor(0x0B, 0x2A, 0x3B), RGBColor(0x2E, 0x8B, 0x57), RGBColor(0x5B, 0x6B, 0x76)
prs = Presentation("slides_tp3.pptx")
doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Pt(64)
    sec.top_margin = sec.bottom_margin = Pt(56)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for nome, tam in (("Title", 24), ("Heading 1", 14)):
    s = doc.styles[nome]
    s.font.name, s.font.size, s.font.bold, s.font.color.rgb = "Cambria", Pt(tam), True, NAVY
    s.element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    s.element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")


def sombra(cell, hexcor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcor)
    tcPr.append(sh)


doc.add_paragraph("Roteiro de apresentação", style="Title")
p = doc.add_paragraph("TP-III · Heurística + Método Exato (TSP / MTZ)\nCauã Fortes, Dyogo Henrique e Maycon Luiz")
p.runs[0].font.color.rgb = MUTED
doc.add_paragraph(
    "Duração prevista: cerca de 15 minutos. O mesmo texto está nas notas do apresentador de slides_tp3.pptx. "
    "Antes de apresentar, registrar no slide 6 a especificação da máquina (processador, núcleos e RAM).")

itens = []
for i, s in enumerate(prs.slides, 1):
    titulo = next(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    nota = s.notes_slide.notes_text_frame.text
    m = re.match(r"\[(.+?) · ~(\d+:\d+)\]\s*(.*)", nota, re.S)
    quem, tempo, texto = (m.group(1), m.group(2), m.group(3)) if m else ("", "", nota)
    itens.append((i, " ".join(titulo.split()), quem, tempo, texto))

doc.add_paragraph("Divisão e tempos", style="Heading 1")
t = doc.add_table(rows=1, cols=4)
t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ["Slide", "Título", "Quem fala", "Tempo"]):
    c.text = h
    sombra(c, "0B2A3B")
    r = c.paragraphs[0].runs[0]
    r.font.bold, r.font.color.rgb = True, RGBColor(255, 255, 255)
for i, titulo, quem, tempo, _ in itens:
    cells = t.add_row().cells
    for c, v in zip(cells, [str(i), (titulo.title().replace(" Iii", " III") if titulo.isupper() else titulo), quem, "~" + tempo]):
        c.text = v
    if i % 2 == 0:
        for c in cells:
            sombra(c, "F4F7F9")
for row in t.rows:
    for c, w in zip(row.cells, [Pt(40), Pt(230), Pt(110), Pt(55)]):
        c.width = w

doc.add_paragraph("Fala por slide", style="Heading 1")
for i, titulo, quem, tempo, texto in itens:
    h = doc.add_paragraph()
    h.paragraph_format.space_before, h.paragraph_format.keep_with_next = Pt(12), True
    r = h.add_run(f"Slide {i} · {(titulo.title().replace(" Iii", " III") if titulo.isupper() else titulo)}")
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Cambria", Pt(12.5), True, NAVY
    if quem:
        m = doc.add_paragraph()
        m.paragraph_format.keep_with_next = True
        rr = m.add_run(f"{quem} · ~{tempo}")
        rr.font.size, rr.font.bold, rr.font.color.rgb = Pt(10), True, GREEN
    b = doc.add_paragraph(texto)
    b.paragraph_format.space_after = Pt(6)
    b.paragraph_format.line_spacing = 1.15
doc.save("roteiro_apresentacao.docx")
print("ok")
