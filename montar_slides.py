"""Gera slides_tp3.pptx (mesmo estilo visual do TP-II) com roteiro nas notas do apresentador.
Requer: pip install python-pptx   (lê resultados.csv)"""
import csv

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY, GREEN, GOLD = RGBColor(0x0B, 0x2A, 0x3B), RGBColor(0x2E, 0x8B, 0x57), RGBColor(0xB7, 0x79, 0x1F)
BG, MUTED, LINE, WHITE = RGBColor(0xF4, 0xF7, 0xF9), RGBColor(0x5B, 0x6B, 0x76), RGBColor(0xD5, 0xDD, 0xE2), RGBColor(255, 255, 255)
RED, BLUE, TEAL = RGBColor(0xC1, 0x12, 0x1F), RGBColor(0x1D, 0x4E, 0x89), RGBColor(0x2A, 0x9D, 0x8F)
COR = {"heuristica": TEAL, "exato": RED, "warm_start": GOLD, "hibrido_FO": BLUE}
NOME = {"heuristica": "Heurística pura", "exato": "Exato puro", "warm_start": "Exato + warm start",
        "hibrido_FO": "Híbrido (WS + F&O)"}
M = Inches(0.6)
W, H = Inches(10), Inches(5.625)

rows = list(csv.DictReader(open("resultados.csv")))
Ns = sorted({int(r["N"]) for r in rows})
res = {(int(r["N"]), r["metodo"]): r for r in rows}

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
contador = [0]


def txt(slide, x, y, w, h, texto, size=14, bold=False, cor=NAVY, fonte="Calibri", align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    linhas = texto if isinstance(texto, list) else [texto]
    for i, l in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = l
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, fonte
        f.color.rgb = cor
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    return s


def novo(titulo, notas, sub=None):
    s = prs.slides.add_slide(BLANK)
    contador[0] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    txt(s, M, Inches(0.35), W - 2 * M, Inches(0.6), titulo, 28, True, NAVY, "Cambria")
    if sub:
        txt(s, M, Inches(0.95), W - 2 * M, Inches(0.3), sub, 13, False, MUTED)
    txt(s, M, Inches(5.2), Inches(5), Inches(0.25), "Trabalho Prático III · Heurística + Método Exato", 10, False, MUTED)
    txt(s, W - M - Inches(0.5), Inches(5.2), Inches(0.5), Inches(0.25), str(contador[0] + 1), 10, False, MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notas
    return s


def card(s, x, y, w, h, titulo, corpo, cor=GREEN, tsize=15, bsize=13):
    box(s, x, y, w, h)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.18), y + Inches(0.18), Inches(0.1), Inches(0.1))
    o.fill.solid(); o.fill.fore_color.rgb = cor; o.line.fill.background(); o.shadow.inherit = False
    txt(s, x + Inches(0.38), y + Inches(0.1), w - Inches(0.5), Inches(0.3), titulo, tsize, True, NAVY, "Cambria")
    txt(s, x + Inches(0.18), y + Inches(0.5), w - Inches(0.36), h - Inches(0.6), corpo, bsize, False, MUTED)


# 1 · capa -----------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
txt(s, M, Inches(0.8), Inches(8), Inches(0.3), "TRABALHO PRÁTICO III  ·  HEURÍSTICA + MÉTODO EXATO", 12, True, RGBColor(0x8F, 0xD3, 0xB0))
txt(s, M, Inches(1.4), Inches(8.6), Inches(1.6), ["Problema do Caixeiro", "Viajante híbrido"], 40, True, WHITE, "Cambria")
txt(s, M, Inches(3.15), Inches(8.2), Inches(0.8),
    "Warm start e fix-and-optimize sobre a formulação MTZ, resolvida com HiGHS", 18, False, RGBColor(0xCF, 0xDD, 0xE6))
txt(s, M, Inches(4.7), Inches(8), Inches(0.3), "Alunos: Cauã Fortes, Dyogo Henrique e Maycon Luiz", 13, False, RGBColor(0xCF, 0xDD, 0xE6))
s.notes_slide.notes_text_frame.text = (
    "[Cauã · ~0:30] Apresentar o tema: no TP-II resolvemos o TSP com PLI (MTZ) e vimos que o exato estoura o tempo na instância difícil. "
    "Neste trabalho combinamos a heurística do TP-I com o exato para tentar empurrar esse limite. Citar os integrantes do grupo.")

# 2 · pergunta -------------------------------------------------------------------------------------
s = novo("O que queremos descobrir", (
    "[Cauã · ~1:00] Retomar o problema: TSP com MTZ, o mesmo do TP-II. Motivação: o exato prova otimalidade mas não escala; a heurística escala mas não prova nada. "
    "Pergunta central: combinar os dois ajuda? Em que tamanho? Mostrar as quatro abordagens que vamos comparar, todas com o mesmo orçamento de tempo (120 s)."),
    "Mesmo problema do TP-II, mesmo orçamento de tempo para todas as abordagens")
w4 = (W - 2 * M - Inches(0.45)) / 4
dados = [("Heurística pura", "Vizinho mais próximo, 2-opt e Or-opt. Rápida, sem prova de otimalidade.", TEAL),
         ("Exato puro", "Modelo MTZ no HiGHS, sem ajuda. Prova o ótimo, mas escala mal.", RED),
         ("Warm start", "A rota da heurística entra como incumbente (limite primal) na largada.", GOLD),
         ("Híbrido", "Warm start, depois fix-and-optimize, depois B&B completo.", BLUE)]
for i, (t, c, cor) in enumerate(dados):
    card(s, M + i * (w4 + Inches(0.15)), Inches(1.6), w4, Inches(2.5), t, c, cor, 15, 14)
txt(s, M, Inches(4.45), W - 2 * M, Inches(0.5),
    "A partir de que N o exato deixa de fechar o gap? O híbrido empurra esse limite?", 16, True, GREEN, "Cambria")

# 3 · instâncias -----------------------------------------------------------------------------------
s = novo("Instâncias geradas", (
    "[Cauã · ~1:00] Gerador parametrizado por N: pontos uniformes em [0,1000]², distância EUC_2D arredondada, a mesma convenção da TSPLIB do TP-II. "
    "Semente documentada: 2026 + N, então qualquer pessoa reproduz. Seis tamanhos, de 10 (instantâneo) a 80 (exato claramente não fecha). "
    "Lembrar as fórmulas: variáveis = n² − 1 e restrições = 2n + (n−1)(n−2); em N=80 o modelo já tem 6.399 variáveis."),
    "Gerador com semente documentada: semente = 2026 + N")
tb = s.shapes.add_table(7, 4, M, Inches(1.5), Inches(5.6), Inches(3.4)).table
cab = ["N (cidades)", "Semente", "Variáveis", "Restrições"]
vals = [cab] + [[str(n), str(2026 + n), f"{n * n - 1:,}".replace(",", "."), f"{2 * n + (n - 1) * (n - 2):,}".replace(",", ".")] for n in Ns]


def preencher(tabela, vals, size=13, cab_cor=NAVY):
    for i, linha in enumerate(vals):
        for j, v in enumerate(linha):
            c = tabela.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = cab_cor if i == 0 else (WHITE if i % 2 else BG)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = v
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            if p.runs:
                f = p.runs[0].font
                f.size, f.name, f.bold = Pt(size), "Calibri", i == 0
                f.color.rgb = WHITE if i == 0 else NAVY


preencher(tb, vals)
card(s, Inches(6.5), Inches(1.5), Inches(2.9), Inches(3.4), "Como o gerador funciona",
     ["Pontos uniformes em [0, 1000]².", "Distância EUC_2D inteira, igual à TSPLIB do TP-II.",
      "Variáveis: n² − 1", "Restrições: 2n + (n−1)(n−2)"], GREEN, 15, 13)

# 4 · heurística -----------------------------------------------------------------------------------
s = novo("Heurística (adaptada do TP-I)", (
    "[Dyogo · ~1:00] A heurística é a do TP-I adaptada para o TSP: constrói uma rota com vizinho mais próximo e melhora com busca local. "
    "2-opt remove duas arestas e reconecta invertendo o trecho; Or-opt move segmentos de 1 a 3 cidades para a melhor posição. "
    "Rodamos a partir de 20 cidades iniciais e ficamos com a melhor. Em todos os tamanhos leva menos de 2 segundos. "
    "Ela entrega a rota para o híbrido e também é o baseline de comparação."),
    "Constrói, melhora e entrega a rota inicial para o solver")
tw = (W - 2 * M - Inches(0.6)) / 3
for i, (t, c) in enumerate([("1 · Construção", "Vizinho mais próximo, partindo de até 20 cidades iniciais diferentes."),
                            ("2 · Busca local", "2-opt (troca de duas arestas) e Or-opt (move segmentos de 1 a 3 cidades) até não melhorar."),
                            ("3 · Seleção", "Fica a melhor das rotas. Em todos os N leva menos de 2 s, sem prova de otimalidade.")]):
    x = M + i * (tw + Inches(0.3))
    card(s, x, Inches(1.5), tw, Inches(2.3), t, c, GREEN, 16, 15)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + tw + Inches(0.04), Inches(2.5), Inches(0.22), Inches(0.3))
        a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background(); a.shadow.inherit = False
txt(s, M, Inches(4.15), W - 2 * M, Inches(0.7),
    "Papel no trabalho: baseline de comparação e fornecedora do incumbente para o solver.", 15, True, NAVY, "Cambria")

# 5 · híbrido --------------------------------------------------------------------------------------
s = novo("O híbrido em três fases", (
    "[Dyogo · ~2:00] Fase 1, warm start: a rota da heurística entra no HiGHS como solução inicial (limite primal). O solver já larga com um bom corte e pode podar nós. "
    "Fase 2, fix-and-optimize, 40% do orçamento: sorteamos uma cidade-centro, pegamos as m mais próximas e liberamos só as arestas que tocam esse grupo (ligando ao grupo, aos 8 vizinhos mais próximos ou à rota atual). "
    "Todo o resto fica fixado na rota atual e o solver resolve esse sub-problema pequeno. A decisão de projeto é a vizinhança espacial: no TSP euclidiano as melhorias úteis acontecem entre cidades próximas. "
    "O tamanho m é adaptativo: começa em 6, cresce em 2 se o sub-problema fecha sem melhora e diminui em 2 se estoura 3 s. "
    "Fase 3: B&B completo com o melhor incumbente, o que também nos dá o limite dual e o gap."),
    "Warm start, fix-and-optimize e B&B completo no mesmo orçamento")
fw = (W - 2 * M - Inches(0.6)) / 3
fases = [("Fase 1 · Warm start", "Rota da heurística entra no HiGHS como incumbente (limite primal)."),
         ("Fase 2 · Fix-and-optimize", "40% do tempo. Libera a vizinhança de m cidades próximas e fixa o resto na rota atual."),
         ("Fase 3 · B&B completo", "Tempo restante com o melhor incumbente. Fornece o limite dual e o gap.")]
for i, (t, c) in enumerate(fases):
    x = M + i * (fw + Inches(0.3))
    card(s, x, Inches(1.4), fw, Inches(1.95), t, c, [TEAL, BLUE, RED][i], 14, 14)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + fw + Inches(0.04), Inches(2.2), Inches(0.22), Inches(0.3))
        a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background(); a.shadow.inherit = False
box(s, M, Inches(3.6), W - 2 * M, Inches(1.35), fill=NAVY, line=None)
txt(s, M + Inches(0.25), Inches(3.72), W - 2 * M - Inches(0.5), Inches(0.3), "Nossa decisão de projeto: vizinhança espacial", 15, True, WHITE, "Cambria")
txt(s, M + Inches(0.25), Inches(4.1), W - 2 * M - Inches(0.5), Inches(0.8),
    "Cidade-centro sorteada + as m mais próximas. Tamanho m adaptativo: começa em 6, +2 se o sub-MIP fecha sem melhora, −2 se estoura 3 s.",
    13, False, RGBColor(0xCF, 0xDD, 0xE6))

# 6 · metodologia ----------------------------------------------------------------------------------
s = novo("Como executamos os testes", (
    "[Maycon · ~1:00] Mesmo modelo MTZ do TP-II, agora montado direto no HiGHS 1.15.1 via highspy, porque precisamos de warm start, fixação de variáveis e do limite dual. "
    "Uma thread por execução. Orçamento de 120 s por abordagem e por N. Rodamos seis tamanhos em paralelo na mesma máquina. "
    "Métrica principal: custo da melhor solução e gap = (custo − melhor limite inferior conhecido) / custo; o limite inferior vem do melhor limite dual entre as abordagens, já que a heurística não tem limite próprio. "
    "[Registrar aqui as especificações da máquina: processador, núcleos e RAM.]"),
    None)
cards = [("Modelagem", "Python + highspy. Modelo MTZ do TP-II, com variáveis binárias e contínuas."),
         ("Solver", "HiGHS 1.15.1, 1 thread, gap relativo 0, semente fixa."),
         ("Orçamento", "120 s por abordagem e por N, igual para as quatro."),
         ("Métricas", "Custo, gap contra o melhor limite inferior conhecido e tempo.")]
cw = (W - 2 * M - Inches(0.3)) / 2
for i, (t, c) in enumerate(cards):
    card(s, M + (i % 2) * (cw + Inches(0.3)), Inches(1.2) + (i // 2) * Inches(1.9), cw, Inches(1.7), t, c, GREEN, 17, 16)

# 7 · tabela de resultados -------------------------------------------------------------------------
s = novo("Resultados por tamanho", (
    "[Maycon · ~1:30] Ler a tabela da esquerda para a direita. Até N=40 todas as abordagens chegam ao mesmo custo e o exato prova o ótimo; o que muda é o tempo (N=40: 60 s no exato puro, 17 s com warm start). "
    "A partir de N=60 ninguém prova o ótimo. O exato puro termina com gap de 12% em N=60 e 29% em N=80, com rotas bem piores. "
    "Heurística e warm start ficam em 1,8% e 6,9%. Em N=80 o híbrido achou 6906 contra 6949, o único ponto onde o fix-and-optimize melhorou algo. "
    "Verde = ótimo provado dentro dos 120 s."),
    "Custo, gap e tempo; em verde, ótimo provado em 120 s")
t = s.shapes.add_table(len(Ns) + 1, 5, M, Inches(1.35), W - 2 * M, Inches(3.5)).table
t.columns[0].width = Inches(0.9)
for j in range(1, 5):
    t.columns[j].width = int((W - 2 * M - Inches(0.9)) / 4)
meto = ["heuristica", "exato", "warm_start", "hibrido_FO"]
v = [["N"] + [NOME[m] for m in meto]] + [[str(n)] + ["-" for _ in meto] for n in Ns]
preencher(t, v, 13)
for i, n in enumerate(Ns, 1):
    for j, m in enumerate(meto, 1):
        r = res[(n, m)]
        c = t.cell(i, j)
        gap = float(r["gap_pct"]) if r["gap_pct"] else 0.0
        prov = r["provado"] == "True"
        tf = c.text_frame
        tf.text = f"{float(r['custo']):.0f}"
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.name = Pt(14), True, "Calibri"
        p.runs[0].font.color.rgb = GREEN if prov else NAVY
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = ("ótimo · " if prov else f"gap {gap:.1f}% · ").replace(".", ",") + f"{float(r['tempo']):.1f} s".replace(".", ",")
        r2.font.size, r2.font.name = Pt(11), "Calibri"
        r2.font.color.rgb = GREEN if prov else MUTED


# 8 / 9 · gráficos nativos -------------------------------------------------------------------------
def grafico(s, campo, titulo_eixo, escala_log=False, y=1.35, h=3.5, fmt="0"):
    cd = CategoryChartData()
    cd.categories = [str(n) for n in Ns]
    for m in meto:
        vals = []
        for n in Ns:
            r = res[(n, m)]
            x = float(r[campo]) if r[campo] else 0.0
            vals.append(max(x, 0.002) if escala_log else x)
        cd.add_series(NOME[m], vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, M, Inches(y), W - 2 * M, Inches(h), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    ch.legend.font.size, ch.legend.font.color.rgb = Pt(12), NAVY
    for ser, m in zip(ch.plots[0].series, meto):
        ser.format.line.color.rgb, ser.format.line.width = COR[m], Pt(2.5)
        ser.smooth = False
        ser.marker.style, ser.marker.size = XL_MARKER_STYLE.CIRCLE, 8
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = COR[m]
        ser.marker.format.line.color.rgb = COR[m]
    va, ca = ch.value_axis, ch.category_axis
    for ax in (va, ca):
        ax.tick_labels.font.size, ax.tick_labels.font.color.rgb = Pt(12), MUTED
        ax.format.line.color.rgb = LINE
    va.major_gridlines.format.line.color.rgb = LINE
    va.has_title, va.axis_title.text_frame.text = True, titulo_eixo
    ca.has_title, ca.axis_title.text_frame.text = True, "N (cidades)"
    for ax in (va, ca):
        f = ax.axis_title.text_frame.paragraphs[0].runs[0].font
        f.size, f.bold, f.color.rgb = Pt(12), False, MUTED
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = fmt, False
    if escala_log:
        sc = va._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
        lb = etree.SubElement(sc, "{http://schemas.openxmlformats.org/drawingml/2006/chart}logBase")
        lb.set("val", "10")
        sc.insert(0, lb)
        va.minimum_scale, va.maximum_scale = 0.001, 1000
    return ch


s = novo("Gap em função do tamanho", (
    "[Maycon · ~1:00] Eixo y é o gap em relação ao melhor limite inferior conhecido para cada N. O exato puro (vermelho) cola em zero até N=40 e dispara em N=60 e N=80. "
    "As outras três quase se sobrepõem: o ganho vem de ter um bom incumbente. O MTZ tem relaxação fraca, então o limite dual quase não sobe em nenhuma das abordagens; "
    "por isso o híbrido melhora a solução mas não consegue fechar o gap."),
    "Gap contra o melhor limite inferior conhecido para cada N")
grafico(s, "gap_pct", "Gap (%)")
s = novo("Tempo em função do tamanho", (
    "[Maycon · ~1:00] Escala logarítmica. A heurística (verde) fica abaixo de 2 s em todos os tamanhos. O exato puro cresce exponencialmente e bate o limite de 120 s a partir de N=60. "
    "O warm start fecha mais rápido em N=30 e N=40. O híbrido com fix-and-optimize gasta pelo menos 40% do orçamento mesmo quando a heurística já é ótima, por isso aparece lá em cima nos tamanhos pequenos: é o custo do fix-and-optimize sem benefício."),
    "Escala logarítmica: o exato cresce de forma exponencial")
grafico(s, "tempo", "Tempo até o fim (s)", escala_log=True, fmt="General")

# 10 · respostas -----------------------------------------------------------------------------------
s = novo("Respostas às perguntas do trabalho", (
    "[Dyogo · ~2:30] 1) O exato deixa de fechar entre N=40, onde prova o ótimo em cerca de 60 s no limite do orçamento, e N=60, com gap de 12%. "
    "2) O híbrido não empurra o limite de fechar o gap: em N=60 e 80 ninguém prova o ótimo. O ganho é na qualidade da solução: gap de 12% para 1,8% em N=60 e de 29% para 6,4% em N=80. O warm start também acelera a prova quando ela é possível (N=40: 17 s contra 60 s). "
    "3) Até N=60 a heurística pura já iguala o híbrido em menos de 1 s; em N=80 o híbrido é 0,6% melhor. "
    "4) O warm start acelerou; o fix-and-optimize quase não fez diferença, porque a heurística 2-opt + Or-opt já é muito forte em instâncias euclidianas deste tamanho e o sub-MIP em MTZ também sofre com a relaxação fraca. "
    "Um F&O mais útil pediria vizinhanças maiores ou uma heurística inicial mais fraca."),
    None)
q = [("Quando o exato deixa de fechar?", "Entre N=40 (ótimo em ~60 s) e N=60 (gap de 12% em 120 s).", RED),
     ("O híbrido empurra o limite?", "Não para fechar o gap, mas sim a qualidade: gap de 12% para 1,8% em N=60 e de 29% para 6,4% em N=80.", BLUE),
     ("A heurística já basta?", "Até N=60 ela iguala o híbrido em menos de 1 s. Em N=80 o híbrido é 0,6% melhor.", TEAL),
     ("A estratégia acelerou o solver?", "O warm start sim (N=40: 17 s contra 60 s). O fix-and-optimize quase não fez diferença.", GOLD)]
qw = (W - 2 * M - Inches(0.25)) / 2
for i, (t, c, cor) in enumerate(q):
    card(s, M + (i % 2) * (qw + Inches(0.25)), Inches(1.15) + (i // 2) * Inches(1.95), qw, Inches(1.8), t, c, cor, 16, 15)

# 11 · conclusões ----------------------------------------------------------------------------------
s = novo("Conclusões", (
    "[Dyogo e Maycon · ~1:30] Três pontos. Um: o warm start é barato e vale a pena, pois acelera a prova de otimalidade e melhora muito a solução quando o exato estoura o tempo. "
    "Dois: o fix-and-optimize só compensa quando a heurística deixa espaço para melhorar; aqui custou 40% do orçamento sem ganho até N=60. "
    "Três: o MTZ limita todo o resto, porque o limite dual é fraco. Ressaltar as limitações: uma instância por N, uma execução por abordagem, seis processos em paralelo na mesma máquina e o gap da heurística usando o limite dual de outra abordagem. "
    "Fechar convidando às perguntas."),
    None)
cc = [("Warm start compensa", "Barato, acelera a prova (N=40: 17 s contra 60 s) e evita soluções ruins quando o tempo acaba."),
      ("F&O depende da heurística", "Com uma heurística já forte, gastou 40% do orçamento sem ganho até N=60."),
      ("O MTZ limita tudo", "A relaxação fraca mantém o limite dual baixo, e nenhum método fecha o gap em N ≥ 60.")]
for i, (t, c) in enumerate(cc):
    card(s, M, Inches(1.2) + i * Inches(1.05), W - 2 * M, Inches(0.95), t, c, [GREEN, GOLD, RED][i], 16, 14)
txt(s, M, Inches(4.45), W - 2 * M, Inches(0.6),
    "Limitações: uma instância por N, uma execução por abordagem, seis processos em paralelo na mesma máquina.", 12, False, MUTED, italic=True)

# 12 · encerramento --------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
txt(s, M, Inches(1.9), W - 2 * M, Inches(0.9), "Obrigado! Dúvidas?", 40, True, WHITE, "Cambria")
txt(s, M, Inches(3.1), W - 2 * M, Inches(0.8), ["TP3: heurística + método exato (TSP / MTZ)",
                                              "Cauã Fortes · Dyogo Henrique · Maycon Luiz"], 16, False, RGBColor(0xCF, 0xDD, 0xE6))
s.notes_slide.notes_text_frame.text = "[Todos · ~2:00] Perguntas. Se perguntarem por que não usamos outra vizinhança: testamos a espacial por fazer sentido para TSP euclidiano; sequências no tour seriam a próxima opção."

prs.save("slides_tp3.pptx")
print("ok", len(prs.slides._sldIdLst), "slides")
